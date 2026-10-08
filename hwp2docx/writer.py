"""문서 모델을 .docx로 쓴다.

python-docx는 패키지 배관(콘텐츠 타입, 관계, 이미지 삽입)만 쓰고,
문단 간격·표 테두리·셀 병합처럼 세밀한 부분은 OOXML을 직접 만든다.
python-docx의 상위 API로는 셀마다 다른 테두리를 줄 수 없기 때문이다.
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Emu

from .docinfo import (
    ALIGN_CENTER,
    ALIGN_DISTRIBUTE,
    ALIGN_DIVIDE,
    ALIGN_JUSTIFY,
    ALIGN_LEFT,
    ALIGN_RIGHT,
    Border,
    BorderFill,
    CharShape,
    DocInfo,
    ParaShape,
)
from .model import Cell, Document, Image, Paragraph, Section, Table
from .reader import hwp_to_dxa, hwp_to_emu

try:  # python-docx 1.x
    from docx.oxml import OxmlElement
except ImportError:  # pragma: no cover
    from docx.oxml.shared import OxmlElement

_ALIGN_TO_OOXML = {
    ALIGN_JUSTIFY: "both",
    ALIGN_LEFT: "left",
    ALIGN_RIGHT: "right",
    ALIGN_CENTER: "center",
    ALIGN_DISTRIBUTE: "distribute",
    ALIGN_DIVIDE: "distribute",
}

#: 한컴 전용 글꼴 -> 어디서나 있는 글꼴. --safe-fonts 일 때만 쓴다.
SAFE_FONT_MAP = {
    "함초롬바탕": "맑은 고딕",
    "함초롬돋움": "맑은 고딕",
    "한컴바탕": "바탕",
    "한컴돋움": "돋움",
    "한컴ソウル": "맑은 고딕",
    "HY헤드라인M": "맑은 고딕",
    "HY견고딕": "맑은 고딕",
    "HY중고딕": "맑은 고딕",
    "HY신명조": "바탕",
    "휴먼명조": "바탕",
    "휴먼고딕": "돋움",
    "새굴림": "굴림",
    "한양신명조": "바탕",
}


def _el(tag: str, **attrs) -> "OxmlElement":
    node = OxmlElement(tag)
    for key, value in attrs.items():
        if value is not None:
            node.set(qn(f"w:{key}"), str(value))
    return node


class DocxWriter:
    def __init__(self, doc: Document, info: DocInfo, safe_fonts: bool = False,
                 font_map: dict[str, str] | None = None, use_layout: bool = True):
        self.doc = doc
        self.info = info
        self.use_layout = use_layout
        self.font_map = dict(SAFE_FONT_MAP) if safe_fonts else {}
        if font_map:
            self.font_map.update(font_map)
        self.docx = DocxDocument()
        self._strip_default_body()

    # -- 공개 --------------------------------------------------------------
    def build(self) -> DocxDocument:
        total = len(self.doc.sections)
        for index, section in enumerate(self.doc.sections):
            self._write_body(section)
            # 마지막 구역의 설정만 body 끝의 sectPr에 들어간다. 그 앞 구역들은
            # 각자 끝에 sectPr을 품은 문단을 두어 구역을 닫는다.
            if index < total - 1:
                self._apply_page(section, self._close_section())
            else:
                self._apply_page(section, self._body_section_properties())
        self._set_default_style()
        return self.docx

    # -- 세로 배치 ----------------------------------------------------------
    def _write_body(self, section: Section) -> None:
        """구역의 문단과 표를 순서대로 내보낸다.

        한글이 PARA_LINE_SEG에 배치 결과를 남겨 둔 문서라면 그 좌표를 써서
        문단 사이 간격을 직접 계산한다. 한글과 Word는 줄 간격·문단 간격을
        쌓는 규칙이 서로 달라서, 문단 모양 값을 그대로 옮기면 문서가
        아래로 밀려 내려간다. 배치 좌표를 쓰면 그 차이가 사라진다.
        """
        body = self.docx.element.body
        use_layout = self.use_layout and any(p.has_layout for p in section.paragraphs)
        if not use_layout:
            for para in section.paragraphs:
                self._write_paragraph(para, body)
            return

        cursor = 0  # 지금까지 채운 높이 (HWPUNIT, 본문 상단 기준)
        paragraphs = section.paragraphs
        for index, para in enumerate(paragraphs):
            if para.tables:
                # 표는 바로 다음 문단 바로 위에 놓인다. 다음 문단의 y에서
                # 표 높이를 빼면 표가 시작하는 지점이 나온다.
                next_top = self._next_top(paragraphs, index)
                for table in para.tables:
                    height = table.total_height or sum(table.row_heights())
                    top = (next_top - height) if next_top is not None else cursor
                    self._spacer(body, top - cursor)
                    self._write_table(table, body)
                    cursor = top + height
                    next_top = None
                if para.runs or para.images:
                    body.append(self._paragraph_element(para, space_before=0))
                    cursor = para.layout_bottom or cursor
                continue

            gap = (para.layout_top - cursor) if para.has_layout else None
            body.append(self._paragraph_element(para, space_before=gap))
            if para.has_layout:
                cursor = para.layout_bottom

        # 본문이 표로 끝나면 Word가 파일을 거부한다.
        if len(body) and body[-1].tag == qn("w:tbl"):
            body.append(self._hairline_paragraph())

    @staticmethod
    def _next_top(paragraphs: list[Paragraph], index: int) -> int | None:
        """index 다음에 오는, 배치 정보가 있는 문단의 시작 y."""
        for para in paragraphs[index + 1 :]:
            if para.has_layout and not para.tables:
                return para.layout_top
        return None

    def _spacer(self, parent, height: int) -> None:
        """정확히 지정한 높이만큼의 빈 문단. Word 표에는 '위 간격'이 없어서
        한글이 표 앞에 두는 빈 문단과 같은 방식으로 자리를 만든다."""
        if height <= 20:
            return
        p = _el("w:p")
        ppr = _el("w:pPr")
        ppr.append(_el("w:spacing", before="0", after="0",
                       line=hwp_to_dxa(height), lineRule="exact"))
        rpr = _el("w:rPr")
        rpr.append(_el("w:sz", val="2"))
        rpr.append(_el("w:szCs", val="2"))
        ppr.append(rpr)
        p.append(ppr)
        parent.append(p)

    def save(self, path: str) -> None:
        self.build().save(path)

    # -- 초기화 ------------------------------------------------------------
    def _strip_default_body(self) -> None:
        """python-docx가 넣어 두는 빈 문단을 지운다."""
        body = self.docx.element.body
        for child in list(body):
            if child.tag == qn("w:p"):
                body.remove(child)
        self._fix_settings()

    @staticmethod
    def _fix_settings_element(settings) -> None:
        """python-docx 기본 템플릿의 w:zoom에는 필수 속성이 빠져 있다."""
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    def _fix_settings(self) -> None:
        try:
            self._fix_settings_element(self.docx.settings.element)
        except Exception:
            pass

    def _set_default_style(self) -> None:
        """Normal 스타일의 여백을 0으로. HWP 값이 그대로 보이게 하려는 것."""
        styles = self.docx.styles.element
        for style in styles.findall(qn("w:style")):
            if style.get(qn("w:styleId")) == "Normal":
                ppr = style.find(qn("w:pPr"))
                if ppr is None:
                    ppr = _el("w:pPr")
                    style.append(ppr)
                spacing = _el("w:spacing", after="0", before="0", line="240", lineRule="auto")
                ppr.append(spacing)
                return

    # -- 구역/페이지 --------------------------------------------------------
    def _close_section(self):
        """구역을 닫는 문단을 넣고 그 sectPr을 돌려준다."""
        para = _el("w:p")
        ppr = _el("w:pPr")
        sect_pr = _el("w:sectPr")
        ppr.append(sect_pr)
        para.append(ppr)
        self.docx.element.body.append(para)
        return sect_pr

    def _body_section_properties(self):
        body = self.docx.element.body
        sect_pr = body.find(qn("w:sectPr"))
        if sect_pr is None:
            sect_pr = _el("w:sectPr")
        else:
            body.remove(sect_pr)
        # sectPr은 반드시 body의 마지막 요소여야 한다.
        body.append(sect_pr)
        return sect_pr

    def _apply_page(self, section: Section, sect_pr) -> None:
        page = section.page

        for child in list(sect_pr):
            if child.tag in (qn("w:pgSz"), qn("w:pgMar")):
                sect_pr.remove(child)

        width, height = hwp_to_dxa(page.width), hwp_to_dxa(page.height)
        pg_sz = _el("w:pgSz", w=width, h=height)
        if page.landscape:
            pg_sz.set(qn("w:orient"), "landscape")
        sect_pr.insert(0, pg_sz)

        # HWP의 위/아래 여백은 머리말·꼬리말 '바깥'이다.
        # Word의 top/bottom은 본문까지의 거리이므로 더해 준다.
        sect_pr.insert(1, _el(
            "w:pgMar",
            top=hwp_to_dxa(page.margin_top + page.header),
            bottom=hwp_to_dxa(page.margin_bottom + page.footer),
            left=hwp_to_dxa(page.margin_left),
            right=hwp_to_dxa(page.margin_right),
            header=hwp_to_dxa(page.margin_top),
            footer=hwp_to_dxa(page.margin_bottom),
            gutter=hwp_to_dxa(page.gutter),
        ))

    # -- 문단 --------------------------------------------------------------
    def _write_paragraph(self, para: Paragraph, parent) -> None:
        # HWP는 표를 문단에 매단다. Word에서는 표가 독립 블록이므로
        # 표를 먼저 내보내고, 남은 글자가 있으면 그 뒤에 문단을 둔다.
        for table in para.tables:
            self._write_table(table, parent)

        if not para.tables or para.runs or para.images:
            parent.append(self._paragraph_element(para))

    def _paragraph_element(self, para: Paragraph, space_before: int | None = None):
        shape = self.info.para_shape(para.para_shape_id)
        p = _el("w:p")
        p.append(self._paragraph_properties(para, shape, space_before))

        for run in para.runs:
            p.append(self._run_element(run))
        for image in para.images:
            p.append(self._image_run(image))
        return p

    def _paragraph_properties(self, para: Paragraph, shape: ParaShape,
                              space_before: int | None = None):
        # w:pPr 하위 요소는 스키마 순서를 지켜야 한다:
        # keepNext -> pageBreakBefore -> widowControl -> pBdr -> shd
        # -> spacing -> ind -> jc -> rPr
        ppr = _el("w:pPr")

        if shape.keep_with_next:
            ppr.append(_el("w:keepNext"))
        if shape.page_break_before or para.page_break_before:
            ppr.append(_el("w:pageBreakBefore"))
        if shape.widow_control:
            ppr.append(_el("w:widowControl"))

        border_fill = self.info.border_fill(shape.border_fill_id) if shape.border_fill_id else None
        if border_fill and any(b.visible for b in
                               (border_fill.left, border_fill.right, border_fill.top, border_fill.bottom)):
            ppr.append(self._borders_element("w:pBdr", border_fill))
        if border_fill and border_fill.fill_color:
            ppr.append(_el("w:shd", val="clear", color="auto", fill=border_fill.fill_color))

        if space_before is not None and para.line_height:
            # 배치 좌표에서 계산한 간격. 줄 높이도 한글이 쓴 값을 그대로 쓴다.
            spacing = {
                "before": max(0, hwp_to_dxa(space_before)),
                "after": 0,
                "line": max(1, hwp_to_dxa(para.line_height)),
                "lineRule": "exact",
            }
        else:
            spacing = {
                "before": hwp_to_dxa(shape.space_before),
                "after": hwp_to_dxa(shape.space_after),
            }
            if shape.line_spacing_type == 0:
                # 글자 크기에 대한 %. Word의 auto 줄 간격은 240이 100%.
                spacing["line"] = max(1, int(round(shape.line_spacing / 100 * 240)))
                spacing["lineRule"] = "auto"
            elif shape.line_spacing_type == 1:
                spacing["line"] = max(1, hwp_to_dxa(shape.line_spacing))
                spacing["lineRule"] = "exact"
            else:
                spacing["line"] = 240
                spacing["lineRule"] = "auto"
        ppr.append(_el("w:spacing", **spacing))

        # 한글의 내어쓰기(음수 들여쓰기)는 첫 줄을 왼쪽 여백에 두고 둘째 줄부터
        # |값|만큼 들여보낸다. Word의 hanging은 반대로 첫 줄을 왼쪽 여백에서
        # 빼낸다. 그래서 Word의 left = 한글 left + |값| 이어야 모양이 같다.
        left = hwp_to_dxa(shape.margin_left)
        right = hwp_to_dxa(shape.margin_right)
        indent = hwp_to_dxa(shape.indent)
        if left or right or indent:
            attrs = {}
            if indent < 0:
                left += -indent
            if left:
                attrs["left"] = left
            if right:
                attrs["right"] = right
            if indent > 0:
                attrs["firstLine"] = indent
            elif indent < 0:
                attrs["hanging"] = -indent
            ppr.append(_el("w:ind", **attrs))

        align = _ALIGN_TO_OOXML.get(shape.align)
        if align:
            ppr.append(_el("w:jc", val=align))

        # 글자가 없는 문단도 줄 높이를 유지하도록 글자 모양을 붙인다.
        if not para.runs:
            rpr = self._run_properties(self.info.char_shape(para.default_char_shape_id))
            ppr.append(rpr)
        return ppr

    # -- 런 ----------------------------------------------------------------
    def _run_element(self, run):
        r = _el("w:r")
        shape = self.info.char_shape(run.char_shape_id)
        r.append(self._run_properties(shape))
        if run.kind == "tab":
            r.append(_el("w:tab"))
        elif run.kind == "break":
            r.append(_el("w:br"))
        else:
            t = _el("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = run.text
            r.append(t)
        return r

    def _resolve_font(self, name: str) -> str:
        return self.font_map.get(name, name)

    def _run_properties(self, shape: CharShape):
        rpr = _el("w:rPr")

        # face_ids는 [한글, 영문, 한자, 일어, 기타, 기호, 사용자] 순
        hangul = self._resolve_font(self.info.font_name(shape.face_ids[0]))
        latin = self._resolve_font(self.info.font_name(shape.face_ids[1])) or hangul
        if hangul or latin:
            fonts = _el("w:rFonts", ascii=latin or hangul, hAnsi=latin or hangul,
                        eastAsia=hangul or latin, cs=latin or hangul)
            rpr.append(fonts)

        # w:rPr 하위 요소도 스키마 순서를 지켜야 한다:
        # rFonts -> b -> bCs -> i -> iCs -> strike -> color -> spacing -> w
        # -> sz -> szCs -> u -> shd -> vertAlign
        if shape.bold:
            rpr.append(_el("w:b"))
            rpr.append(_el("w:bCs"))
        if shape.italic:
            rpr.append(_el("w:i"))
            rpr.append(_el("w:iCs"))
        if shape.strike:
            rpr.append(_el("w:strike"))
        if shape.color:
            rpr.append(_el("w:color", val=shape.color))

        # 자간: 한글은 %, Word는 1/20 pt
        if shape.spacing:
            pt = shape.base_size / 100.0 * shape.spacing / 100.0
            rpr.append(_el("w:spacing", val=int(round(pt * 20))))
        # 장평
        if shape.ratio and shape.ratio != 100:
            rpr.append(_el("w:w", val=max(1, min(600, shape.ratio))))

        size = shape.half_points
        rpr.append(_el("w:sz", val=size))
        rpr.append(_el("w:szCs", val=size))

        if shape.underline:
            rpr.append(_el("w:u", val="single", color=shape.underline_color or None))
        if shape.shade_color:
            rpr.append(_el("w:shd", val="clear", color="auto", fill=shape.shade_color))
        if shape.superscript:
            rpr.append(_el("w:vertAlign", val="superscript"))
        elif shape.subscript:
            rpr.append(_el("w:vertAlign", val="subscript"))
        return rpr

    # -- 그림 --------------------------------------------------------------
    def _image_run(self, image: Image):
        """python-docx로 그림을 넣고, 만들어진 run 요소를 가져온다."""
        try:
            holder = self.docx.add_paragraph()
            run = holder.add_run()
            width = hwp_to_emu(image.width) or None
            height = hwp_to_emu(image.height) or None
            run.add_picture(io.BytesIO(image.data),
                            width=Emu(width) if width else None,
                            height=Emu(height) if height else None)
            element = run._r
            holder._p.remove(element)
            self.docx.element.body.remove(holder._p)
            return element
        except Exception:
            self.doc.warn("그림 하나를 Word에 넣지 못했습니다.")
            return _el("w:r")

    # -- 표 ----------------------------------------------------------------
    def _write_table(self, table: Table, parent) -> None:
        grid = table.grid()
        widths = [hwp_to_dxa(w) for w in grid]
        total = sum(widths)
        if total <= 0:
            return

        # Word는 두 표가 맞붙으면 하나로 합쳐 버린다. 사이에 높이가 거의 없는
        # 문단을 끼워 막는다. 한 줄짜리 빈 문단을 쓰면 표마다 간격이 쌓인다.
        if len(parent) and parent[-1].tag == qn("w:tbl"):
            parent.append(self._hairline_paragraph())

        tbl = _el("w:tbl")
        tbl.append(self._table_properties(table, total))

        tbl_grid = _el("w:tblGrid")
        for w in widths:
            tbl_grid.append(_el("w:gridCol", w=w))
        tbl.append(tbl_grid)

        occupied: set[tuple[int, int]] = set()
        by_position = {(c.row, c.col): c for c in table.cells}
        row_heights = table.row_heights()

        for row_index in range(table.rows):
            tr = _el("w:tr")
            height = row_heights[row_index] if row_index < len(row_heights) else 0
            if height:
                trpr = _el("w:trPr")
                trpr.append(_el("w:trHeight", val=hwp_to_dxa(height), hRule="atLeast"))
                tr.append(trpr)

            col = 0
            while col < table.cols:
                cell = by_position.get((row_index, col))
                if cell is not None:
                    tr.append(self._cell_element(cell, widths, table))
                    for r in range(cell.row_span):
                        for c in range(cell.col_span):
                            occupied.add((row_index + r, col + c))
                    col += cell.col_span
                elif (row_index, col) in occupied:
                    # 세로 병합으로 이어지는 칸
                    tr.append(self._continuation_cell(widths[col] if col < len(widths) else 0))
                    col += 1
                else:
                    tr.append(self._empty_cell(widths[col] if col < len(widths) else 0))
                    col += 1
            tbl.append(tr)

        parent.append(tbl)

    @staticmethod
    def _hairline_paragraph():
        """높이를 거의 차지하지 않는 빈 문단. 표 구분용."""
        p = _el("w:p")
        ppr = _el("w:pPr")
        ppr.append(_el("w:spacing", before="0", after="0", line="1", lineRule="exact"))
        rpr = _el("w:rPr")
        rpr.append(_el("w:sz", val="2"))
        rpr.append(_el("w:szCs", val="2"))
        ppr.append(rpr)
        p.append(ppr)
        return p

    def _table_properties(self, table: Table, total_width: int):
        tbl_pr = _el("w:tblPr")
        # tblPr 하위 요소는 스키마가 순서를 강제한다:
        # tblW -> jc -> tblCellSpacing -> tblInd -> tblBorders -> shd
        # -> tblLayout -> tblCellMar -> tblLook
        tbl_pr.append(_el("w:tblW", w=total_width, type="dxa"))
        if table.cell_spacing:
            tbl_pr.append(_el("w:tblCellSpacing", w=hwp_to_dxa(table.cell_spacing), type="dxa"))
        # 표 테두리 위치 = 여백선 + tblInd - 표 기본 왼쪽 셀 여백.
        # 한글은 테두리를 여백선에 맞추므로 둘 다 0으로 둔다.
        # (안쪽 여백은 셀마다 tcMar로 따로 준다.)
        tbl_pr.append(_el("w:tblInd", w=0, type="dxa"))

        border_fill = self.info.border_fill(table.border_fill_id) if table.border_fill_id else None
        if border_fill:
            tbl_pr.append(self._borders_element("w:tblBorders", border_fill, inside=True))

        tbl_pr.append(_el("w:tblLayout", type="fixed"))

        # 셀 안쪽 여백의 기본값. 셀마다 다시 덮어쓴다.
        margins = _el("w:tblCellMar")
        for side, value in (("top", 0), ("left", 0), ("bottom", 0), ("right", 0)):
            margins.append(_el(f"w:{side}", w=value, type="dxa"))
        tbl_pr.append(margins)
        return tbl_pr

    def _borders_element(self, tag: str, fill: BorderFill, inside: bool = False):
        node = _el(tag)
        pairs = [("top", fill.top), ("left", fill.left), ("bottom", fill.bottom), ("right", fill.right)]
        if inside:
            pairs += [("insideH", fill.top), ("insideV", fill.left)]
        for name, border in pairs:
            node.append(self._border_element(f"w:{name}", border))
        return node

    @staticmethod
    def _border_element(tag: str, border: Border):
        if not border.visible:
            return _el(tag, val="nil")
        return _el(tag, val=border.style, sz=border.size, space="0",
                   color=border.color or "auto")

    def _cell_element(self, cell: Cell, widths: list[int], table: Table):
        tc = _el("w:tc")
        tc_pr = _el("w:tcPr")

        span_width = sum(widths[cell.col : cell.col + cell.col_span]) or (
            widths[cell.col] if cell.col < len(widths) else 0)
        tc_pr.append(_el("w:tcW", w=span_width, type="dxa"))
        if cell.col_span > 1:
            tc_pr.append(_el("w:gridSpan", val=cell.col_span))
        if cell.row_span > 1:
            tc_pr.append(_el("w:vMerge", val="restart"))

        fill_id = cell.border_fill_id or table.border_fill_id
        border_fill = self.info.border_fill(fill_id) if fill_id else None
        if border_fill:
            tc_pr.append(self._borders_element("w:tcBorders", border_fill))
            if border_fill.fill_color:
                tc_pr.append(_el("w:shd", val="clear", color="auto", fill=border_fill.fill_color))

        # w:tcPr 순서: tcW -> gridSpan -> vMerge -> tcBorders -> shd
        # -> tcMar -> vAlign
        left, right, top, bottom = cell.margins
        margins = _el("w:tcMar")
        for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            margins.append(_el(f"w:{name}", w=hwp_to_dxa(value), type="dxa"))
        tc_pr.append(margins)
        tc_pr.append(_el("w:vAlign",
                         val={0: "top", 1: "center", 2: "bottom"}.get(cell.vertical_align, "top")))
        tc.append(tc_pr)

        wrote = False
        for para in cell.paragraphs:
            for nested in para.tables:
                self._write_table(nested, tc)
                wrote = True
            if not para.tables or para.runs:
                # 셀 안에서는 한글도 문단 간격을 쌓지 않는다. 줄 높이만 맞춘다.
                before = 0 if (self.use_layout and para.has_layout) else None
                tc.append(self._paragraph_element(para, space_before=before))
                wrote = True
        if not wrote:
            tc.append(_el("w:p"))
        # Word는 셀이 문단으로 끝나야 한다.
        if tc[-1].tag != qn("w:p"):
            tc.append(_el("w:p"))
        return tc

    @staticmethod
    def _continuation_cell(width: int):
        tc = _el("w:tc")
        tc_pr = _el("w:tcPr")
        tc_pr.append(_el("w:tcW", w=width, type="dxa"))
        tc_pr.append(_el("w:vMerge"))
        tc.append(tc_pr)
        tc.append(_el("w:p"))
        return tc

    @staticmethod
    def _empty_cell(width: int):
        tc = _el("w:tc")
        tc_pr = _el("w:tcPr")
        tc_pr.append(_el("w:tcW", w=width, type="dxa"))
        tc.append(tc_pr)
        tc.append(_el("w:p"))
        return tc


def write_docx(doc: Document, info: DocInfo, path: str, safe_fonts: bool = False,
               font_map: dict[str, str] | None = None, use_layout: bool = True) -> None:
    DocxWriter(doc, info, safe_fonts=safe_fonts, font_map=font_map,
               use_layout=use_layout).save(path)
